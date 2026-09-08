from __future__ import annotations

import fnmatch
import hashlib
import json
import os
import re
import sqlite3
import urllib.error
import urllib.parse
import urllib.request
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable

TEXT_SUFFIXES = {
    ".md", ".txt", ".rst", ".json", ".yaml", ".yml",
    ".py", ".js", ".jsx", ".ts", ".tsx", ".java", ".go",
    ".rs", ".c", ".h", ".cpp", ".hpp", ".sh", ".sql",
}
SUPPORTED_SUFFIXES = TEXT_SUFFIXES | {".docx", ".xlsx"}
SENSITIVE_NAMES = {
    ".env", ".env.local", ".env.production",
    "id_rsa", "id_ed25519",
}
SENSITIVE_SUFFIXES = {".pem", ".key", ".p12", ".pfx"}
SENSITIVE_PARTS = {".ssh", ".aws", ".kube", ".git"}

DEFAULT_MAX_FILE_BYTES = 8 * 1024 * 1024
DEFAULT_MAX_TEXT_CHARS = 40_000
DEFAULT_MAX_ITEMS = 500
DEFAULT_XLSX_ROWS = 200
DEFAULT_XLSX_COLS = 30


def utc_now() -> str:
    return datetime.now(timezone.utc).isoformat()


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as fh:
        while chunk := fh.read(1024 * 1024):
            digest.update(chunk)
    return digest.hexdigest()


def is_sensitive_path(path: Path, project_root: Path) -> bool:
    try:
        relative = path.relative_to(project_root)
    except ValueError:
        return True
    if path.name in SENSITIVE_NAMES or path.suffix.lower() in SENSITIVE_SUFFIXES:
        return True
    return any(part in SENSITIVE_PARTS for part in relative.parts)


def is_ignored(relative: str, patterns: Iterable[str]) -> bool:
    path = relative.replace("\\", "/")
    name = Path(relative).name
    for pattern in patterns:
        normalized = pattern.replace("\\", "/").rstrip("/")
        if fnmatch.fnmatch(path, normalized) or fnmatch.fnmatch(name, normalized):
            return True
        if path == normalized or path.startswith(normalized + "/"):
            return True
    return False


def expand_declared_path(project_root: Path, raw: str) -> Iterable[Path]:
    has_glob = any(ch in raw for ch in "*?[")
    candidates = list(project_root.glob(raw)) if has_glob else [project_root / raw]
    for candidate in candidates:
        if candidate.is_file():
            yield candidate
        elif candidate.is_dir():
            yield from (p for p in candidate.rglob("*") if p.is_file())


def iter_analysis_files(project_root: Path, manifest: dict[str, Any]) -> Iterable[tuple[Path, str]]:
    analysis_cfg = manifest.get("analysis") or {}
    include = set(analysis_cfg.get("include") or ["documents", "tests", "outputs"])
    paths_cfg = manifest.get("paths") or {}
    ignore = list(manifest.get("ignore") or [])
    seen: set[str] = set()

    for category in ("documents", "tests", "outputs", "code"):
        if category not in include:
            continue
        for raw in paths_cfg.get(category) or []:
            for path in expand_declared_path(project_root, raw):
                try:
                    relative = path.relative_to(project_root).as_posix()
                except ValueError:
                    continue
                if relative in seen or is_ignored(relative, ignore):
                    continue
                seen.add(relative)
                yield path, category


def classify_role(relative: str, category: str) -> str:
    text = relative.lower()
    rules = [
        ("requirement", ["需求", "requirement", "prd", "需求说明", "需求规格"]),
        ("test_result", ["测试", "test", "验收", "验证", "评测", "qa"]),
        ("handoff", ["交接", "handoff", "移交"]),
        ("report", ["周报", "日报", "月报", "报告", "report"]),
        ("design", ["设计", "架构", "architecture", "design", "方案"]),
        ("interface", ["接口", "api", "interface"]),
        ("deployment", ["部署", "deploy", "docker", "compose"]),
        ("decision", ["决策", "decision", "会议纪要", "meeting"]),
    ]
    for role, keywords in rules:
        if any(keyword in text for keyword in keywords):
            return role
    return {
        "documents": "document",
        "tests": "test_result",
        "outputs": "output",
        "code": "code",
    }.get(category, "document")


def _limit_text(parts: Iterable[str], limit: int) -> str:
    output: list[str] = []
    length = 0
    for raw in parts:
        value = str(raw).strip()
        if not value:
            continue
        remaining = limit - length
        if remaining <= 0:
            break
        value = value[:remaining]
        output.append(value)
        length += len(value) + 1
    return "\n".join(output)


def parse_text(path: Path, max_chars: int) -> tuple[str, dict[str, Any]]:
    text = path.read_text(encoding="utf-8", errors="replace")
    return text[:max_chars], {"parser": "text", "truncated": len(text) > max_chars}


def parse_docx(path: Path, max_chars: int) -> tuple[str, dict[str, Any]]:
    from docx import Document

    document = Document(str(path))
    chunks: list[str] = []
    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            chunks.append(paragraph.text)
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                chunks.append(" | ".join(cells))
    text = _limit_text(chunks, max_chars)
    return text, {
        "parser": "docx",
        "paragraphs": len(document.paragraphs),
        "tables": len(document.tables),
        "truncated": sum(len(x) for x in chunks) > max_chars,
    }


def parse_xlsx(path: Path, max_chars: int, max_rows: int, max_cols: int) -> tuple[str, dict[str, Any]]:
    from openpyxl import load_workbook

    workbook = load_workbook(str(path), read_only=True, data_only=True)
    chunks: list[str] = []
    sheets: list[dict[str, Any]] = []
    try:
        for worksheet in workbook.worksheets:
            chunks.append(f"[Sheet] {worksheet.title}")
            row_count = 0
            for row in worksheet.iter_rows(values_only=True):
                row_count += 1
                values = []
                for value in row[:max_cols]:
                    if value is None:
                        values.append("")
                    else:
                        values.append(str(value).replace("\n", " ").strip())
                if any(values):
                    chunks.append("\t".join(values))
                if row_count >= max_rows:
                    break
            sheets.append({"name": worksheet.title, "sampled_rows": row_count})
    finally:
        workbook.close()

    text = _limit_text(chunks, max_chars)
    return text, {
        "parser": "xlsx",
        "sheets": sheets,
        "sample_limits": {"rows_per_sheet": max_rows, "columns": max_cols},
        "truncated": len(text) >= max_chars,
    }


def parse_file(path: Path, max_chars: int, xlsx_rows: int, xlsx_cols: int) -> tuple[str, dict[str, Any]]:
    suffix = path.suffix.lower()
    if suffix in TEXT_SUFFIXES:
        return parse_text(path, max_chars)
    if suffix == ".docx":
        return parse_docx(path, max_chars)
    if suffix == ".xlsx":
        return parse_xlsx(path, max_chars, xlsx_rows, xlsx_cols)
    raise ValueError(f"unsupported suffix: {suffix}")


def _compact_line(line: str, max_len: int = 300) -> str:
    return re.sub(r"\s+", " ", line).strip()[:max_len]


def deterministic_analysis(text: str, role: str, relative: str) -> dict[str, Any]:
    lines = [_compact_line(line) for line in text.splitlines()]
    lines = [line for line in lines if line]
    summary = "；".join(lines[:3])[:600]

    keyword_rules = [
        ("blocker", ["阻塞", "失败", "错误", "无法", "异常", "blocker", "failed", "error"]),
        ("requirement", ["需求", "要求", "必须", "应当", "requirement", "must", "should"]),
        ("decision", ["决定", "确定", "采用", "调整为", "decision"]),
        ("task", ["待办", "下一步", "进行中", "todo", "next step"]),
        ("test", ["测试", "通过", "不通过", "passed", "failed", "test"]),
        ("deployment", ["部署", "上线", "发布", "deploy", "release"]),
        ("milestone", ["里程碑", "阶段完成", "验收", "milestone"]),
    ]
    facts: list[dict[str, str]] = []
    seen: set[tuple[str, str]] = set()
    for line in lines[:500]:
        lower = line.lower()
        for fact_type, keywords in keyword_rules:
            if any(keyword in lower for keyword in keywords):
                key = (fact_type, line)
                if key not in seen:
                    seen.add(key)
                    facts.append({"type": fact_type, "text": line})
                break
        if len(facts) >= 30:
            break

    return {
        "source": "deterministic",
        "document_role": role,
        "summary": summary or f"{relative}：未提取到可读文本",
        "facts": facts,
        "mentioned_modules": [],
    }


def endpoint_is_local(base_url: str) -> bool:
    try:
        hostname = urllib.parse.urlparse(base_url).hostname
    except ValueError:
        return False
    return hostname in {"localhost", "127.0.0.1", "::1", "host.docker.internal"}


def analyze_with_openai_compatible(
    *,
    text: str,
    role: str,
    relative: str,
    base_url: str,
    model: str,
    api_key: str,
    timeout: int,
) -> dict[str, Any]:
    endpoint = base_url.rstrip("/") + "/chat/completions"
    system = (
        "你是研发项目资料分析器。只根据输入文件内容提取事实，不补充外部知识。"
        "输出严格 JSON，字段必须为 document_role、summary、facts、mentioned_modules。"
        "facts 为数组，每项包含 type 和 text；type 只能是 requirement、decision、task、test、"
        "blocker、milestone、deployment、note。不要输出源码正文或大段原文。"
    )
    user = (
        f"文件: {relative}\n预分类: {role}\n"
        "请提取对项目进度管理有价值的事实。若信息不足就少输出，不要猜测。\n\n"
        f"{text}"
    )
    payload = json.dumps(
        {
            "model": model,
            "temperature": 0,
            "response_format": {"type": "json_object"},
            "messages": [
                {"role": "system", "content": system},
                {"role": "user", "content": user},
            ],
        },
        ensure_ascii=False,
    ).encode("utf-8")
    headers = {"Content-Type": "application/json"}
    if api_key:
        headers["Authorization"] = f"Bearer {api_key}"
    request = urllib.request.Request(endpoint, data=payload, headers=headers, method="POST")
    with urllib.request.urlopen(request, timeout=timeout) as response:
        data = json.loads(response.read().decode("utf-8"))
    content = data["choices"][0]["message"]["content"].strip()
    if content.startswith("```"):
        content = re.sub(r"^```(?:json)?\s*|\s*```$", "", content, flags=re.IGNORECASE)
    parsed = json.loads(content)
    facts = []
    for item in parsed.get("facts") or []:
        if not isinstance(item, dict):
            continue
        fact_type = str(item.get("type") or "note")
        text_value = _compact_line(str(item.get("text") or ""), 500)
        if text_value:
            facts.append({"type": fact_type, "text": text_value})
    return {
        "source": "llm",
        "document_role": str(parsed.get("document_role") or role),
        "summary": _compact_line(str(parsed.get("summary") or ""), 1000),
        "facts": facts[:50],
        "mentioned_modules": [
            _compact_line(str(item), 120)
            for item in (parsed.get("mentioned_modules") or [])[:30]
            if str(item).strip()
        ],
    }


class AnalysisCache:
    def __init__(self, path: Path) -> None:
        path.parent.mkdir(parents=True, exist_ok=True)
        self.path = path
        with self.connect() as conn:
            conn.execute(
                """
                CREATE TABLE IF NOT EXISTS file_analysis (
                    path TEXT PRIMARY KEY,
                    sha256 TEXT NOT NULL,
                    role TEXT NOT NULL,
                    parser TEXT,
                    analyzed_at TEXT NOT NULL,
                    result_json TEXT NOT NULL
                )
                """
            )

    def connect(self) -> sqlite3.Connection:
        conn = sqlite3.connect(self.path)
        conn.row_factory = sqlite3.Row
        return conn

    def get(self, relative: str, sha256: str) -> dict[str, Any] | None:
        with self.connect() as conn:
            row = conn.execute(
                "SELECT result_json FROM file_analysis WHERE path=? AND sha256=?",
                (relative, sha256),
            ).fetchone()
        return json.loads(row["result_json"]) if row else None

    def put(self, relative: str, sha256: str, role: str, parser: str, result: dict[str, Any]) -> None:
        with self.connect() as conn:
            conn.execute(
                """
                INSERT INTO file_analysis(path, sha256, role, parser, analyzed_at, result_json)
                VALUES (?, ?, ?, ?, ?, ?)
                ON CONFLICT(path) DO UPDATE SET
                    sha256=excluded.sha256,
                    role=excluded.role,
                    parser=excluded.parser,
                    analyzed_at=excluded.analyzed_at,
                    result_json=excluded.result_json
                """,
                (
                    relative,
                    sha256,
                    role,
                    parser,
                    utc_now(),
                    json.dumps(result, ensure_ascii=False),
                ),
            )

    def prune(self, seen_paths: set[str]) -> int:
        with self.connect() as conn:
            rows = conn.execute("SELECT path FROM file_analysis").fetchall()
            stale = [row["path"] for row in rows if row["path"] not in seen_paths]
            if stale:
                conn.executemany("DELETE FROM file_analysis WHERE path=?", [(path,) for path in stale])
        return len(stale)


def analyze_project_files(
    project_root: Path,
    manifest: dict[str, Any],
    state_root: Path,
) -> dict[str, Any]:
    security = manifest.get("security") or {}
    mode = security.get("mode", "metadata_only")
    analysis_cfg = manifest.get("analysis") or {}

    if mode != "local_analysis":
        return {
            "enabled": False,
            "mode": mode,
            "reason": "正文分析仅在 security.mode=local_analysis 时启用",
        }
    if analysis_cfg.get("enabled", True) is False:
        return {"enabled": False, "mode": mode, "reason": "analysis.enabled=false"}

    max_file_bytes = int(analysis_cfg.get("max_file_bytes", DEFAULT_MAX_FILE_BYTES))
    max_text_chars = int(analysis_cfg.get("max_text_chars", DEFAULT_MAX_TEXT_CHARS))
    max_items = int(analysis_cfg.get("max_items", DEFAULT_MAX_ITEMS))
    xlsx_rows = int(analysis_cfg.get("xlsx_rows_per_sheet", DEFAULT_XLSX_ROWS))
    xlsx_cols = int(analysis_cfg.get("xlsx_columns", DEFAULT_XLSX_COLS))

    use_llm = bool(analysis_cfg.get("use_llm", False))
    base_url = os.getenv("LOCAL_LLM_BASE_URL", "").strip()
    model = os.getenv("LOCAL_LLM_MODEL", "").strip()
    api_key = os.getenv("LOCAL_LLM_API_KEY", "")
    timeout = int(os.getenv("LOCAL_LLM_TIMEOUT", "60"))
    allow_remote = os.getenv("ALLOW_REMOTE_ANALYSIS_ENDPOINT", "false").lower() == "true"
    llm_ready = use_llm and bool(base_url and model) and (allow_remote or endpoint_is_local(base_url))

    project_id = str((manifest.get("project") or {}).get("id") or project_root.name)
    safe_project_id = re.sub(r"[^a-zA-Z0-9._-]+", "_", project_id)
    cache = AnalysisCache(state_root / f"{safe_project_id}.sqlite3")

    stats = Counter()
    role_counts = Counter()
    items: list[dict[str, Any]] = []
    seen_paths: set[str] = set()

    for path, category in iter_analysis_files(project_root, manifest):
        if len(items) >= max_items:
            stats["truncated"] += 1
            break
        try:
            relative = path.relative_to(project_root).as_posix()
            size = path.stat().st_size
        except OSError:
            stats["io_error"] += 1
            continue

        seen_paths.add(relative)
        if is_sensitive_path(path, project_root):
            stats["sensitive_skipped"] += 1
            continue
        if path.suffix.lower() not in SUPPORTED_SUFFIXES:
            stats["unsupported"] += 1
            continue
        if size > max_file_bytes:
            stats["too_large"] += 1
            continue

        role = classify_role(relative, category)
        role_counts[role] += 1
        try:
            digest = sha256_file(path)
        except OSError:
            stats["io_error"] += 1
            continue

        cached = cache.get(relative, digest)
        if cached is not None:
            stats["cached"] += 1
            items.append(cached)
            continue

        try:
            text, parser_meta = parse_file(path, max_text_chars, xlsx_rows, xlsx_cols)
        except Exception as exc:
            stats["parse_error"] += 1
            items.append(
                {
                    "path": relative,
                    "sha256": digest,
                    "role": role,
                    "status": "parse_error",
                    "error": type(exc).__name__,
                }
            )
            continue

        analysis_result: dict[str, Any]
        if llm_ready and text.strip():
            try:
                analysis_result = analyze_with_openai_compatible(
                    text=text,
                    role=role,
                    relative=relative,
                    base_url=base_url,
                    model=model,
                    api_key=api_key,
                    timeout=timeout,
                )
                stats["llm_analyzed"] += 1
            except (urllib.error.URLError, TimeoutError, KeyError, ValueError, json.JSONDecodeError):
                analysis_result = deterministic_analysis(text, role, relative)
                analysis_result["llm_fallback"] = True
                stats["llm_fallback"] += 1
        else:
            analysis_result = deterministic_analysis(text, role, relative)
            stats["deterministic_analyzed"] += 1

        item = {
            "path": relative,
            "sha256": digest,
            "role": role,
            "status": "analyzed",
            "parser": parser_meta,
            "analysis": analysis_result,
        }
        cache.put(relative, digest, role, str(parser_meta.get("parser") or ""), item)
        items.append(item)
        stats["new_or_changed"] += 1

    stats["stale_removed"] = cache.prune(seen_paths)

    facts: list[dict[str, Any]] = []
    for item in items:
        analysis = item.get("analysis") or {}
        for fact in analysis.get("facts") or []:
            facts.append(
                {
                    "path": item.get("path"),
                    "role": item.get("role"),
                    "type": fact.get("type"),
                    "text": fact.get("text"),
                }
            )
            if len(facts) >= 200:
                break
        if len(facts) >= 200:
            break

    return {
        "enabled": True,
        "mode": mode,
        "llm": {
            "requested": use_llm,
            "ready": llm_ready,
            "model": model or None,
            "remote_endpoint_allowed": allow_remote,
        },
        "stats": dict(stats),
        "project_memory": {
            "document_role_counts": dict(role_counts),
            "facts": facts,
        },
        "items": items,
    }
