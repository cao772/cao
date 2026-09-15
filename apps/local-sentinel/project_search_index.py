from __future__ import annotations

import csv
import io
from collections import Counter
from pathlib import Path
from typing import Any

from document_pipeline import TEXT_SUFFIXES, is_sensitive_path, iter_analysis_files

DEFAULT_MAX_FILE_BYTES = 8 * 1024 * 1024
DEFAULT_MAX_ITEMS = 300
DEFAULT_MAX_CHARS_PER_FILE = 6000
DEFAULT_MAX_SEGMENTS_PER_FILE = 60


def _clean_text(value: Any, limit: int = 900) -> str:
    return " ".join(str(value or "").replace("\x00", " ").split())[:limit]


def _append_segment(
    segments: list[dict[str, Any]],
    *,
    text: Any,
    locator: str,
    location_type: str,
    max_segments: int,
    max_chars: int,
    char_count: list[int],
    extra: dict[str, Any] | None = None,
) -> bool:
    if len(segments) >= max_segments or char_count[0] >= max_chars:
        return False
    cleaned = _clean_text(text, min(900, max_chars - char_count[0]))
    if not cleaned:
        return True
    item = {
        "text": cleaned,
        "locator": locator,
        "location_type": location_type,
    }
    if extra:
        item.update(extra)
    segments.append(item)
    char_count[0] += len(cleaned)
    return len(segments) < max_segments and char_count[0] < max_chars


def _index_text(path: Path, max_segments: int, max_chars: int) -> list[dict[str, Any]]:
    text = path.read_text(encoding="utf-8", errors="replace")
    segments: list[dict[str, Any]] = []
    chars = [0]
    for line_number, line in enumerate(text.splitlines(), start=1):
        if not _append_segment(
            segments,
            text=line,
            locator=f"第{line_number}行",
            location_type="line",
            max_segments=max_segments,
            max_chars=max_chars,
            char_count=chars,
            extra={"line": line_number},
        ):
            break
    return segments


def _index_docx(path: Path, max_segments: int, max_chars: int) -> list[dict[str, Any]]:
    from docx import Document

    document = Document(str(path))
    segments: list[dict[str, Any]] = []
    chars = [0]
    for index, paragraph in enumerate(document.paragraphs, start=1):
        if not _append_segment(
            segments,
            text=paragraph.text,
            locator=f"第{index}段",
            location_type="paragraph",
            max_segments=max_segments,
            max_chars=max_chars,
            char_count=chars,
            extra={"paragraph": index},
        ):
            return segments
    for table_index, table in enumerate(document.tables, start=1):
        for row_index, row in enumerate(table.rows, start=1):
            text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if not _append_segment(
                segments,
                text=text,
                locator=f"表{table_index}第{row_index}行",
                location_type="table_row",
                max_segments=max_segments,
                max_chars=max_chars,
                char_count=chars,
                extra={"table": table_index, "row": row_index},
            ):
                return segments
    return segments


def _index_xlsx(path: Path, max_segments: int, max_chars: int) -> list[dict[str, Any]]:
    from openpyxl import load_workbook

    workbook = load_workbook(str(path), read_only=True, data_only=True)
    segments: list[dict[str, Any]] = []
    chars = [0]
    try:
        for worksheet in workbook.worksheets:
            for row_index, row in enumerate(worksheet.iter_rows(values_only=True), start=1):
                values = [_clean_text(value, 220) for value in row[:40] if value is not None]
                text = " | ".join(value for value in values if value)
                if not _append_segment(
                    segments,
                    text=text,
                    locator=f"{worksheet.title}!第{row_index}行",
                    location_type="sheet_row",
                    max_segments=max_segments,
                    max_chars=max_chars,
                    char_count=chars,
                    extra={"sheet": worksheet.title, "row": row_index},
                ):
                    return segments
    finally:
        workbook.close()
    return segments


def _decode_csv(raw: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def _index_csv(path: Path, max_segments: int, max_chars: int) -> list[dict[str, Any]]:
    reader = csv.reader(io.StringIO(_decode_csv(path.read_bytes())))
    segments: list[dict[str, Any]] = []
    chars = [0]
    for row_index, row in enumerate(reader, start=1):
        text = " | ".join(_clean_text(value, 240) for value in row[:50] if str(value).strip())
        if not _append_segment(
            segments,
            text=text,
            locator=f"第{row_index}行",
            location_type="row",
            max_segments=max_segments,
            max_chars=max_chars,
            char_count=chars,
            extra={"row": row_index},
        ):
            break
    return segments


def _index_pdf(path: Path, max_segments: int, max_chars: int) -> list[dict[str, Any]]:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    segments: list[dict[str, Any]] = []
    chars = [0]
    for page_number, page in enumerate(reader.pages[:80], start=1):
        text = page.extract_text() or ""
        paragraphs = [part.strip() for part in text.splitlines() if part.strip()]
        page_text = " ".join(paragraphs)
        if not _append_segment(
            segments,
            text=page_text,
            locator=f"第{page_number}页",
            location_type="page",
            max_segments=max_segments,
            max_chars=max_chars,
            char_count=chars,
            extra={"page": page_number},
        ):
            break
    return segments


def _index_file(path: Path, max_segments: int, max_chars: int) -> list[dict[str, Any]]:
    suffix = path.suffix.lower()
    if suffix in TEXT_SUFFIXES:
        return _index_text(path, max_segments, max_chars)
    if suffix == ".docx":
        return _index_docx(path, max_segments, max_chars)
    if suffix == ".xlsx":
        return _index_xlsx(path, max_segments, max_chars)
    if suffix == ".csv":
        return _index_csv(path, max_segments, max_chars)
    if suffix == ".pdf":
        return _index_pdf(path, max_segments, max_chars)
    return []


def build_project_search_index(project_root: Path, manifest: dict[str, Any]) -> dict[str, Any]:
    security = manifest.get("security") or {}
    analysis = manifest.get("analysis") or {}
    if security.get("mode", "metadata_only") != "local_analysis":
        return {
            "enabled": False,
            "reason": "内容检索仅在 local_analysis 模式启用",
            "items": [],
        }
    if analysis.get("search_index_enabled", True) is False:
        return {"enabled": False, "reason": "analysis.search_index_enabled=false", "items": []}

    max_file_bytes = int(analysis.get("max_file_bytes", DEFAULT_MAX_FILE_BYTES))
    max_items = int(analysis.get("search_index_max_items", DEFAULT_MAX_ITEMS))
    max_chars = int(analysis.get("search_index_chars_per_file", DEFAULT_MAX_CHARS_PER_FILE))
    max_segments = int(analysis.get("search_index_segments_per_file", DEFAULT_MAX_SEGMENTS_PER_FILE))

    stats = Counter()
    items: list[dict[str, Any]] = []
    for path, category in iter_analysis_files(project_root, manifest):
        if len(items) >= max_items:
            stats["truncated"] += 1
            break
        if category == "code":
            stats["code_skipped"] += 1
            continue
        try:
            relative = path.relative_to(project_root).as_posix()
            size = path.stat().st_size
        except (OSError, ValueError):
            stats["io_error"] += 1
            continue
        if is_sensitive_path(path, project_root):
            stats["sensitive_skipped"] += 1
            continue
        if size > max_file_bytes:
            stats["too_large"] += 1
            continue
        try:
            segments = _index_file(path, max_segments, max_chars)
        except Exception:
            stats["parse_error"] += 1
            continue
        if not segments:
            stats["empty_or_unsupported"] += 1
            continue
        items.append(
            {
                "path": relative,
                "category": category,
                "suffix": path.suffix.lower(),
                "segments": segments,
                "segment_count": len(segments),
                "indexed_chars": sum(len(str(segment.get("text") or "")) for segment in segments),
            }
        )
        stats["indexed"] += 1

    return {
        "enabled": True,
        "mode": "bounded_content_segments",
        "limits": {
            "max_items": max_items,
            "chars_per_file": max_chars,
            "segments_per_file": max_segments,
        },
        "stats": dict(stats),
        "items": items,
    }
